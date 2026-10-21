import re
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

from app.domain.models import Row
from app.domain.text import fold

# Feldname → erlaubte Kopfzeilen (gefaltet, klein, ohne Leerzeichen)
ALIASES: dict[str, tuple[str, ...]] = {
    "order": ("bestellung", "bestellnummer"),
    "company": ("firma",),
    "name": ("teilnehmer", "name"),
    "email": ("email", "e-mail", "mail"),
    "street": ("strasse", "str."),
    "house": ("hausnummer", "hausnr", "nr"),
    "stair": ("stiege",),
    "door": ("tuer", "tur"),
    "zip": ("plz",),
    "city": ("ort",),
    "country": ("land",),
    "phone": ("telefon", "tel"),
    "notes": ("notizen", "notiz"),
    "special": ("spezialwuensche", "spezialwunsch", "spezialwunsche"),
}


class InputError(ValueError):
    pass


@dataclass
class Sheet:
    rows: list[Row]
    header_row: int
    columns: dict[str, int]  # Feld → 1-basierte Spalte
    title: str | None
    has_special_column: bool


def _key(value) -> str:
    return re.sub(r"\s+", "", fold(str(value or "")).lower())


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_header(ws, max_scan: int = 10) -> int:
    for r in range(1, min(ws.max_row, max_scan) + 1):
        keys = {_key(c.value) for c in ws[r]}
        if "bestellung" in keys and "teilnehmer" in keys:
            return r
    raise InputError("Keine Kopfzeile mit „Bestellung“ und „Teilnehmer“ gefunden.")


def map_columns(ws, header_row: int) -> dict[str, int]:
    columns: dict[str, int] = {}
    for c in ws[header_row]:
        k = _key(c.value)
        for field, names in ALIASES.items():
            if field not in columns and k in names:
                columns[field] = c.column
    return columns


def read_sheet(data: bytes) -> Sheet:
    try:
        wb = load_workbook(BytesIO(data))
    except Exception as exc:  # openpyxl wirft verschiedene Typen
        raise InputError("Die Datei ist keine gültige Excel-Datei (.xlsx).") from exc
    ws = wb.worksheets[0]
    header = find_header(ws)
    columns = map_columns(ws, header)
    title = None
    if header > 1:
        title = next((_cell(c.value) for r in range(1, header) for c in ws[r] if c.value), None)
    rows: list[Row] = []
    for r in range(header + 1, ws.max_row + 1):
        values = {f: _cell(ws.cell(r, col).value) for f, col in columns.items()}
        if not values.get("order"):
            continue
        rows.append(Row(index=r, **values))
    return Sheet(rows=rows, header_row=header, columns=columns, title=title, has_special_column="special" in columns)


def order_key(order: str) -> tuple[int, str]:
    digits = re.sub(r"\D", "", order)
    return (int(digits) if digits else 0, order)
