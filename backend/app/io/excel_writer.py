from copy import copy
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.derived import SORTIMENT_HEADER, sortiment_rows
from app.domain.models import MEALS, Meal, Row, RowResult, Run
from app.domain.rules import Catalog
from app.io.excel_reader import find_header, map_columns, order_key

YELLOW = PatternFill("solid", fgColor="FFFF00")
MEAL_FILL = {
    Meal.F: PatternFill("solid", fgColor="92D050"),
    Meal.M: PatternFill("solid", fgColor="FFC000"),
    Meal.A: PatternFill("solid", fgColor="00B0F0"),
}
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)


def _save(wb: Workbook, path: Path) -> None:
    buf = BytesIO()
    wb.save(buf)
    path.write_bytes(buf.getvalue())


def relevant_pairs(run: Run) -> list[tuple[Row, RowResult]]:
    pairs = [(row, res) for row, res in zip(run.rows, run.results) if res.relevant]
    return sorted(pairs, key=lambda p: order_key(p[0].order))


def write_gesamt(run: Run, upload: bytes, path: Path) -> None:
    wb = load_workbook(BytesIO(upload))
    ws = wb.worksheets[0]
    header = find_header(ws)
    shift = 0
    if header == 1:
        ws.insert_rows(1)
        header, shift = 2, 1
    columns = map_columns(ws, header)
    col = columns.get("special")
    if col is None:
        col = ws.max_column + 1
        src = ws.cell(header, col - 1)
        cell = ws.cell(header, col, "Spezialwünsche")
        cell.font, cell.fill, cell.border, cell.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)
        ws.column_dimensions[get_column_letter(col)].width = 40
    title = ws.cell(1, 1, f"{run.delivery.label} {run.boxes} Boxen")
    title.fill, title.font = YELLOW, BOLD
    wish = {res.index: res.wish_text for res in run.results}
    for row in run.rows:
        cell = ws.cell(row.index + shift, col)
        text = wish.get(row.index, "")
        cell.value = text or None
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    _save(wb, path)


LABEL_HEADER = ["Teilnehmer", "Strasse", "Hausnummer", "Stiege", "Tür", "PLZ", "Spezialwünsche"]


def write_labels(run: Run, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Etikettenliste"
    ws.append(LABEL_HEADER)
    for c in ws[1]:
        c.font = BOLD
    for row, res in relevant_pairs(run):
        ws.append([row.name, row.street, row.house, row.stair, row.door, row.zip, res.wish_text])
    for letter, width in zip("ABCDEFG", (28, 28, 12, 8, 8, 8, 50)):
        ws.column_dimensions[letter].width = width
    _save(wb, path)


def summary_cells(run: Run, catalog: Catalog) -> dict[Meal, str]:
    out: dict[Meal, str] = {}
    for meal in MEALS:
        t = run.totals.meals[meal]
        parts = [f"{n}x Ohne {catalog.by_id[rid].label if rid in catalog.by_id else rid}"
                 for rid, n in t.without.items() if n]
        if parts:
            out[meal] = ", ".join(parts)
    return out


def _style_like(cell, src) -> None:
    cell.font, cell.fill, cell.border, cell.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)


def write_special(run: Run, catalog: Catalog, path: Path, template: Path | None) -> None:
    if template and template.exists():
        wb = load_workbook(template)
        ws = wb.worksheets[0]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Tabelle1"
        for c, v in zip("ABCDEF", ("Datum:", None, "Boxenanzahl:", "Früh", "Mittag", "Abend")):
            ws[f"{c}1"] = v
        for c, v in zip("ABC", ("Bestellung", "Teilnehmer", "Spezialwünsche")):
            ws[f"{c}2"] = v
        for r in (1, 2):
            for c in "ABCDEF":
                cell = ws[f"{c}{r}"]
                cell.font, cell.border = BOLD, BOX
                cell.fill = YELLOW
        for meal, c in zip(MEALS, "DEF"):
            ws[f"{c}1"].fill = ws[f"{c}2"].fill = MEAL_FILL[meal]
        for letter, width in zip("ABCDEF", (15, 33, 42, 12, 20, 20)):
            ws.column_dimensions[letter].width = width
        ws.merge_cells("A1:B1")

    t = run.totals
    ws["A1"] = f"Datum: {run.delivery.label}"
    ws["C1"] = f"Boxenanzahl: {run.boxes}"
    for meal, c in zip(MEALS, "DEF"):
        ws[f"{c}2"] = t.meals[meal].total
    addons = [r for r in catalog.addons() if t.addons.get(r.id)]
    for i, rule in enumerate(addons):
        cell = ws.cell(2, 7 + i, rule.label)
        cell.font = BOLD
        ws.column_dimensions[get_column_letter(7 + i)].width = max(12, len(rule.label) + 2)

    pairs = relevant_pairs(run)
    first = 3
    for n, (row, res) in enumerate(pairs):
        r = first + n
        ws.cell(r, 1, row.order)
        ws.cell(r, 2, row.name)
        ws.cell(r, 3, res.wish_text).alignment = Alignment(wrap_text=True, vertical="top")
        for meal, c in zip(MEALS, (4, 5, 6)):
            ws.cell(r, c, res.meal_cells.get(meal) or None)
        for i, rule in enumerate(addons):
            cell = ws.cell(r, 7 + i, res.addons.get(rule.id))
            _style_like(cell, ws.cell(r, 6))
            cell.fill = PatternFill()
    last = first + len(pairs) - 1
    total_row = last + 1
    for meal, c in zip(MEALS, (4, 5, 6)):
        text = summary_cells(run, catalog).get(meal)
        cell = ws.cell(total_row, c, text)
        if text:
            cell.fill, cell.font = MEAL_FILL[meal], BOLD
    for i in range(len(addons)):
        letter = get_column_letter(7 + i)
        cell = ws.cell(total_row, 7 + i, f"=SUM({letter}{first}:{letter}{max(last, first)})")
        cell.font = BOLD
    _save(wb, path)


def write_amounts(run: Run, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sortiment neu"
    ws.append([f"Sortiment neu · {run.delivery.weekday} {run.delivery.label}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(SORTIMENT_HEADER + ["", "Kombi-Portionen", "Check"])
    for c in ws[2]:
        if c.value:
            c.font, c.border = BOLD, BOX
    rows = sortiment_rows(run.totals, run.delivery.weekday, run.menu)
    order = (Meal.A, Meal.F, Meal.M)
    for i, (values, meal) in enumerate(zip(rows, order)):
        r = 3 + i
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, int(v) if v.isdigit() else (v or None))
            cell.border = BOX
            if c in (5, 6) and v == "":
                cell.fill = RED
        ws.cell(r, 1).fill = MEAL_FILL[meal]
        ws.cell(r, 8, run.totals.meals[meal].combo).border = BOX
        check = ws.cell(r, 9, f"=C{r}-D{r}-N(E{r})-N(F{r})+H{r}")
        check.border = BOX
        check.fill = GREEN if run.totals.meals[meal].check == 0 else RED
    ws["I6"] = "Check muss 0 sein"
    ws["I6"].font = Font(bold=True, color="C00000")
    for letter, width in zip("ABCDEFGHI", (14, 14, 15, 15, 11, 12, 3, 16, 10)):
        ws.column_dimensions[letter].width = width
    _save(wb, path)
