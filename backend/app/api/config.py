import re
from datetime import date
from io import BytesIO

from fastapi import APIRouter, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel

from app.domain.engine import analyze_row
from app.domain.models import Match, ReviewItem, Row, Rule
from app.domain.rules import Catalog, compile_rule, rules_from_keywords
from app.io import storage

router = APIRouter(prefix="/api/config", tags=["config"])


class RuleTest(BaseModel):
    text: str


class RuleTestResult(BaseModel):
    matches: list[Match]
    review: list[ReviewItem]


def _validate(rules: list[Rule]) -> None:
    ids = [r.id for r in rules]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, "Jede Regel braucht eine eindeutige ID.")
    for r in rules:
        if not r.patterns and not r.case_sensitive:
            raise HTTPException(422, f"Regel „{r.label}“ hat kein Suchmuster.")
        try:
            compile_rule(r)
        except re.error as exc:
            raise HTTPException(422, f"Suchmuster in „{r.label}“ ist ungültig: {exc}") from exc


@router.get("/rules", response_model=list[Rule])
def get_rules() -> list[Rule]:
    return storage.load_rules()


@router.put("/rules", response_model=list[Rule])
def put_rules(rules: list[Rule]) -> list[Rule]:
    rules = [r.model_copy(update={"patterns": [p for p in r.patterns if p.strip()],
                                  "case_sensitive": [c for c in r.case_sensitive if c.strip()]}) for r in rules]
    _validate(rules)
    storage.save_rules(rules)
    return rules


@router.post("/rules/import", response_model=list[Rule])
async def import_rules(file: UploadFile = File(...)) -> list[Rule]:
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(415, "Bitte die Suchwörter als .xlsx hochladen.")
    try:
        ws = load_workbook(BytesIO(await file.read(2 * 1024 * 1024))).worksheets[0]
    except Exception as exc:
        raise HTTPException(422, "Die Datei ist keine gültige Excel-Datei.") from exc
    words = [str(v).strip() for (v,) in ws.iter_rows(min_col=1, max_col=1, values_only=True) if v]
    if words and words[0].lower() == "suchwort":
        words = words[1:]
    rules = rules_from_keywords(words, storage.load_rules())
    storage.save_rules(rules)
    return rules


@router.post("/rules/test", response_model=RuleTestResult)
def test_rules(body: RuleTest) -> RuleTestResult:
    result = analyze_row(Row(index=0, order="0", special=body.text), Catalog(storage.load_rules()), date.today())
    return RuleTestResult(matches=result.matches, review=result.review)


@router.get("/mail-templates")
def get_mail_templates() -> dict:
    return storage.load_mail_templates()


@router.put("/mail-templates")
def put_mail_templates(value: dict) -> dict:
    if not isinstance(value.get("recipients"), list):
        raise HTTPException(422, "Vorlagen brauchen eine Liste „recipients“.")
    storage.save_mail_templates(value)
    return value
