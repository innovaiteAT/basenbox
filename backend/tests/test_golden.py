"""Golden Test gegen das echte Beispiel vom 08.06. (Montag, KW24).

Die Datei enthält Kundendaten und liegt nur lokal (gitignored), daher wird der Test sonst übersprungen.
Geprüft werden nur Summen, keine Zeileninhalte.
"""
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.domain.models import Meal, RuleKind
from app.service import create_run, generate_files, recompute

EXAMPLE = Path(__file__).resolve().parents[2] / "data" / "input" / "example" / "KW24" / "08.06. Montag"
SOURCE = EXAMPLE / "08.06. Gesamt.xlsx"
MENU = {"nuts": [Meal.M], "zk": [Meal.A], "raisins": []}

pytestmark = pytest.mark.skipif(not SOURCE.exists(), reason="Beispieldatei nicht vorhanden")


@pytest.fixture(scope="module")
def run():
    from app.domain.rules import Catalog, default_rules

    data = SOURCE.read_bytes()
    base = create_run(data, SOURCE.name, date(2026, 6, 7), delivery_date=date(2026, 6, 8))
    return recompute(base.model_copy(update={"menu": MENU}), Catalog(default_rules()), {})


def test_totals(run):
    t = run.totals
    assert run.delivery.weekday == "Montag" and run.delivery.iso_week == 24
    assert t.boxes == 61
    assert (t.meals[Meal.F].total, t.meals[Meal.M].total, t.meals[Meal.A].total) == (62, 60, 62)
    assert t.meals[Meal.M].without == {"nuts": 5}
    assert t.meals[Meal.A].without == {"zk": 8}
    # Die Handdatei zeigt 20: dort steht eine Kreatin-1 in einer Zeile ohne Kreatin-Wunsch (Fehler in der Handdatei).
    assert t.addons["kreatin"] == 19
    assert len(t.relevant_rows) == 29
    assert all(m.check == 0 for m in t.meals.values())


def test_review_and_irrelevant_rows(run):
    review_rows = [r for r in run.results if r.review]
    assert len(review_rows) == 1
    assert all(i.reason == "Wunsch ohne passende Regel" for i in review_rows[0].review)
    only_raisins = [r for r in run.results if {m.rule_id for m in r.matches} == {"raisins"}]
    assert only_raisins and not any(r.relevant for r in only_raisins)


def test_kreatin_from_notes_only(run):
    from_notes = [m for r in run.results for m in r.matches
                  if m.kind == RuleKind.addon and m.source_column.value == "notes"]
    assert from_notes, "Kreatin nur in Notizen muss erkannt werden"


def test_files(run, tmp_path):
    from app.domain.rules import Catalog, default_rules

    template = EXAMPLE.parents[2] / "templates" / "Spezialwünsche_Template.xlsx"
    generate_files(run, Catalog(default_rules()), SOURCE.read_bytes(), tmp_path, template)
    ws = load_workbook(tmp_path / "08.06. Spezialwünsche.xlsx").active
    assert [ws.cell(2, c).value for c in (4, 5, 6)] == [62, 60, 62]
    assert ws["E32"].value == "5x Ohne Nüsse" and ws["F32"].value == "8x Ohne Z&K"
    assert ws["G2"].value == "Kreatin" and ws["G32"].value == "=SUM(G3:G31)"
    assert load_workbook(tmp_path / "08.06. Etikettenliste.xlsx").active.max_row == 30
    assert load_workbook(tmp_path / "08.06. Gesamt.xlsx").active["A1"].value == "08.06. 61 Boxen"
