from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.domain.models import Meal
from app.io.excel_reader import read_sheet
from app.service import create_run, generate_files, recompute
from tests.conftest import HEADER, make_xlsx

TODAY = date(2026, 6, 7)
ROWS = [
    {"Notizen": "Stock: 2", "Spezialwünsche": "OM; 2xAbend"},
    {"Notizen": "Kreatin (am 8.6.)"},
    {"Spezialwünsche": "ohne Nüsse"},
    {"Notizen": "Rabatt-Aktion"},
]


def test_header_in_row_two():
    sheet = read_sheet(make_xlsx(ROWS))
    assert sheet.header_row == 2 and sheet.title == "08.06. 3 Boxen" and len(sheet.rows) == 4


def test_header_in_row_one_and_aliases():
    header = [h.replace("Straße", "Strasse").replace("Notizen", "Notiz") for h in HEADER]
    sheet = read_sheet(make_xlsx([{"Notiz": "x", "Strasse": "Weg"}], title=None, header=header))
    assert sheet.header_row == 1 and sheet.title is None
    assert sheet.rows[0].notes == "x" and sheet.rows[0].street == "Weg"


def test_missing_special_column(catalog):
    data = make_xlsx(ROWS, header=HEADER[:-1])
    run = recompute(create_run(data, "08.06. Gesamt.xlsx", TODAY), catalog, {})
    assert not run.has_special_column
    assert "no_special" in [w.code for w in run.warnings]


def test_run_from_filename_and_box_warning(catalog):
    run = recompute(create_run(make_xlsx(ROWS), "08.06. Gesamt.xlsx", TODAY), catalog, {})
    assert run.delivery.date == date(2026, 6, 8) and run.boxes == 4 and run.title_boxes == 3
    assert run.warnings[0].code == "boxes"


def test_generate_files(tmp_path, catalog):
    # ohne Spalte Spezialwünsche: alle Wünsche stehen in den Notizen, die Spalte wird angelegt
    rows = [{"Notizen": "Stock: 2 - OM; 2xAbend"}, ROWS[1], {"Notizen": "ohne Nüsse"}, ROWS[3]]
    data = make_xlsx(rows, header=HEADER[:-1])
    run = create_run(data, "08.06. Gesamt.xlsx", TODAY)
    run = recompute(run.model_copy(update={"menu": {"nuts": [Meal.M], "zk": [Meal.A]}}), catalog, {})
    files = generate_files(run, catalog, data, tmp_path, None)
    assert sorted(files) == sorted(p.name for p in tmp_path.iterdir())

    gesamt = load_workbook(tmp_path / "08.06. Gesamt.xlsx").active
    assert gesamt["A1"].value == "08.06. 4 Boxen"
    assert gesamt.cell(2, 14).value == "Spezialwünsche"
    assert gesamt.cell(3, 14).value == "OM; 2xAbend" and gesamt.cell(3, 14).font.b
    assert gesamt.cell(3, 13).value == "Stock: 2 - OM; 2xAbend"  # Notizen bleiben
    assert gesamt.cell(6, 14).value is None

    labels = load_workbook(tmp_path / "08.06. Etikettenliste.xlsx").active
    assert [c.value for c in labels[1]] == ["Teilnehmer", "Strasse", "Hausnummer", "Stiege", "Tür", "PLZ",
                                             "Spezialwünsche"]
    assert labels.max_row == 4 and not labels.merged_cells.ranges

    special = load_workbook(tmp_path / "08.06. Spezialwünsche.xlsx").active
    assert special["A1"].value == "Datum: 08.06." and special["C1"].value == "Boxenanzahl: 4"
    assert [special.cell(2, c).value for c in range(4, 8)] == [4, 3, 5, "Kreatin"]
    assert special["E6"].value == "1x Ohne Nüsse" and special["G6"].value == "=SUM(G3:G5)"
    assert special["D1"].fill.fgColor.rgb.endswith("92D050")

    amounts = load_workbook(tmp_path / "08.06. genaue Menge.xlsx").active
    assert [amounts.cell(r, 1).value for r in (3, 4, 5)] == ["Abend", "Frühstück", "Mittag"]
    assert (tmp_path / "08.06. Spezialwünsche.pdf").read_bytes().startswith(b"%PDF")


def test_generate_with_template(tmp_path, catalog):
    from openpyxl import Workbook

    tpl = Workbook()
    tpl.active["D1"] = "Früh "
    tpl_path = tmp_path / "tpl.xlsx"
    tpl.save(tpl_path)
    out = tmp_path / "out"
    data = make_xlsx(ROWS)
    run = recompute(create_run(data, "08.06. Gesamt.xlsx", TODAY), catalog, {})
    generate_files(run, catalog, data, out, tpl_path)
    ws = load_workbook(BytesIO((out / "08.06. Spezialwünsche.xlsx").read_bytes())).active
    assert ws["D1"].value == "Früh " and ws["A1"].value == "Datum: 08.06."
